# excel_rules.py
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill

RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def col_letter(df, column_name):
    """Find a column's Excel letter (A, B, C...) by its header name, not a hardcoded position."""
    idx = df.columns.get_loc(column_name) + 1  # openpyxl is 1-indexed
    return get_column_letter(idx)


def add_dropdown(ws, df, column_name, choices, max_row=1000):
    col = col_letter(df, column_name)
    formula = '"' + ",".join(choices) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = f"Value must be one of: {', '.join(choices)}"
    dv.errorTitle = "Invalid entry"
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{max_row}")


def add_numeric_highlight(ws, df, column_name, max_row=1000):
    col = col_letter(df, column_name)
    formula = f'AND({col}2<>"", NOT(ISNUMBER(--{col}2)))'
    ws.conditional_formatting.add(f"{col}2:{col}{max_row}", FormulaRule(formula=[formula], fill=RED_FILL))


def add_date_order_highlight(ws, df, start_col_name, end_col_name, max_row=1000):
    start_col = col_letter(df, start_col_name)
    end_col = col_letter(df, end_col_name)
    formula = f'AND({end_col}2<>"", {start_col}2<>"", {end_col}2<{start_col}2)'
    ws.conditional_formatting.add(f"{end_col}2:{end_col}{max_row}", FormulaRule(formula=[formula], fill=RED_FILL))

def add_date_order_validation(ws, df, start_col_name, end_col_name, max_row=1000):
    start_col = col_letter(df, start_col_name)
    end_col = col_letter(df, end_col_name)
    formula = f'OR({start_col}2="",{end_col}2>={start_col}2)'
    dv = DataValidation(type="custom", formula1=formula, allow_blank=True, showErrorMessage=True)
    dv.error = f"{end_col_name} cannot be earlier than {start_col_name}"
    dv.errorTitle = "Invalid date order"
    ws.add_data_validation(dv)
    dv.add(f"{end_col}2:{end_col}{max_row}")